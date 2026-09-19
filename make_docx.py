"""
Render the preprint manuscript to .docx, with figures and the anchor table.

bioRxiv accepts Word directly, and Word is the more useful of the two here because
the author list, funding and competing-interest statements still have to be filled in
by a human before anything is posted.

The converter handles the subset of Markdown the manuscript actually uses, and three
things a naive line-by-line converter gets wrong: wrapped lines have to be joined into
one paragraph (otherwise every hard wrap becomes a paragraph break), <sup>/<sub> have
to become real super/subscript runs rather than literal tags, and pipe tables have to
become Word tables rather than rows of pipes.
"""

import json
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "preprint", "manuscript.md")
OUT = os.path.join(HERE, "preprint", "Suitcordance_preprint.docx")
FIGS = [("Figure 1", "out/figures/fig1_trajectories.png", 6.4),
        ("Figure 2", "out/figures/fig2_calibration.png", 3.9),
        ("Figure 3", "out/figures/fig3_experiments.png", 6.4)]

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|<sup>.+?</sup>|<sub>.+?</sub>)")


def add_runs(par, text):
    """Bold, italic, code and super/subscript spans; everything else plain."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif piece.startswith("<sup>") and piece.endswith("</sup>"):
            par.add_run(piece[5:-6]).font.superscript = True
        elif piece.startswith("<sub>") and piece.endswith("</sub>"):
            par.add_run(piece[5:-6]).font.subscript = True
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def blocks(lines):
    """Group the manuscript's lines into blocks a Word document can carry."""
    out, buf = [], []

    def flush():
        if buf:
            out.append(("para", " ".join(buf)))
            buf.clear()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped or stripped == "---":
            flush()
            i += 1
            continue

        if line.startswith("    ") and stripped:          # indented code / formulae
            flush()
            code = []
            while i < len(lines) and (lines[i].startswith("    ") or not lines[i].strip()):
                if lines[i].strip():
                    code.append(lines[i].rstrip()[4:])
                elif code:
                    break
                i += 1
            out.append(("code", code))
            continue

        if stripped.startswith("|") and stripped.endswith("|"):   # pipe table
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            out.append(("table", rows))
            continue

        m = re.match(r"^(#{1,3})\s+(.*)", stripped)
        if m:
            flush()
            out.append((f"h{len(m.group(1))}", m.group(2)))
            i += 1
            continue

        for prefix, kind in ((">", "quote"), ("- ", "bullet")):
            if stripped.startswith(prefix):
                flush()
                chunk = [stripped[len(prefix):].strip()]
                i += 1
                while i < len(lines):
                    nxt = lines[i].rstrip()
                    if not nxt.strip():
                        break
                    if nxt.strip().startswith(prefix) or re.match(r"^(#{1,3})\s", nxt.strip()) \
                            or nxt.strip().startswith("|") or re.match(r"^\d+\.\s", nxt.strip()):
                        break
                    chunk.append(nxt.strip())
                    i += 1
                out.append((kind, " ".join(chunk)))
                break
        else:
            if re.match(r"^\d+\.\s", stripped):
                flush()
                chunk = [re.sub(r"^\d+\.\s", "", stripped)]
                i += 1
                while i < len(lines) and lines[i].strip() and \
                        not re.match(r"^\d+\.\s", lines[i].strip()) and \
                        not lines[i].strip().startswith(("#", "|", "-", ">")):
                    chunk.append(lines[i].strip())
                    i += 1
                out.append(("number", " ".join(chunk)))
                continue
            buf.append(stripped)
            i += 1

    flush()
    return out


def write_table(doc, rows, font=Pt(9)):
    if not rows:
        return
    t = doc.add_table(rows=1, cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    for c, h in enumerate(rows[0]):
        cell = t.rows[0].cells[c]
        cell.text = ""
        add_runs(cell.paragraphs[0], h)
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.size = font
    for row in rows[1:]:
        cells = t.add_row().cells
        for c, v in enumerate(row[:len(rows[0])]):
            cells[c].text = ""
            add_runs(cells[c].paragraphs[0], v)
            for r in cells[c].paragraphs[0].runs:
                r.font.size = font


def anchor_table(doc):
    rows = json.load(open(os.path.join(HERE, "out", "anchors.json"), encoding="utf-8"))
    bed = {"coronary": "Coronary", "sfa": "Femoropopliteal", "btk": "Below the knee",
           "carotid": "Carotid", "renal": "Renal", "iliac": "Iliac"}
    p = doc.add_paragraph()
    add_runs(p, "**Table 1. The anchor set, with provenance.** Endpoint constructs are "
                "given as the source words them. Trust weight encodes audited evidence "
                "quality. Rows marked (excluded) are not in the fit and are shown at "
                "their corrected value.")
    data = [["Anchor", "Bed", "Observed", "Predicted", "Construct", "Trigger",
             "Unit", "Window (d)", "Trust"]]
    for r in rows:
        name = r["name"] + ("" if r["included"] else " (excluded)")
        data.append([name, bed.get(r["bed"], r["bed"]),
                     f'{r["observed"]*100:.1f}%', f'{r["predicted"]*100:.1f}%',
                     r["construct"], r["trigger"], r["unit"],
                     str(r["window_days"]), f'{r["weight"]:.2f}'])
    write_table(doc, data, font=Pt(7.5))

    doc.add_paragraph()
    p = doc.add_paragraph()
    add_runs(p, "**Table 1 (continued). Primary citations, and what changed.**")
    for r in rows:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(4)
        run = q.add_run(r["name"] + ". ")
        run.bold = True
        run.font.size = Pt(8.5)
        run2 = q.add_run(r["citation"] + (f' doi:{r["doi"]}' if r["doi"] else ""))
        run2.font.size = Pt(8.5)
        run3 = q.add_run(" — " + r["source"])
        run3.font.size = Pt(8.5)
        run3.italic = True


def main():
    lines = open(SRC, encoding="utf-8").read().splitlines()
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.3

    for kind, payload in blocks(lines):
        if kind == "h1":
            h = doc.add_heading("", level=0)
            add_runs(h, payload)
        elif kind in ("h2", "h3"):
            h = doc.add_heading("", level=int(kind[1]) - 1)
            add_runs(h, payload)
        elif kind == "code":
            for ln in payload:
                p = doc.add_paragraph()
                r = p.add_run(ln)
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
                p.paragraph_format.space_after = Pt(1)
        elif kind == "table":
            write_table(doc, payload)
            doc.add_paragraph()
        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, payload)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xA8, 0x33, 0x2A)
                r.font.size = Pt(9.5)
        elif kind == "bullet":
            add_runs(doc.add_paragraph(style="List Bullet"), payload)
        elif kind == "number":
            add_runs(doc.add_paragraph(style="List Number"), payload)
        else:
            add_runs(doc.add_paragraph(), payload)

    doc.add_page_break()
    h = doc.add_heading("", level=1)
    add_runs(h, "Figures")
    for name, path, width in FIGS:
        full = os.path.join(HERE, path.replace("/", os.sep))
        if not os.path.exists(full):
            raise SystemExit(f"missing figure: {full}")
        doc.add_picture(full, width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        r = cap.add_run(f"{name}. Legend in the text.")
        r.font.size = Pt(9)
        r.italic = True
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()
    h = doc.add_heading("", level=1)
    add_runs(h, "Table 1")
    anchor_table(doc)

    doc.save(OUT)
    print("written", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    main()
